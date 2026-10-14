"""Build Gear_Rebalance.xlsx (detailed) and Gear_Tier_List.xlsx (simple, for friends).

    python tools/build_gear_sheets.py

Inputs: tools/wiki_gear.json (parse_gear_wiki.py), tools/gear_stats.json (gear_stats.py),
tools/artifact_values.json (artifact_values.py), tools/game_en.json (item names).
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
TOOLS = ROOT / "tools"
wiki = json.loads((TOOLS / "wiki_gear.json").read_text(encoding="utf-8"))
stats = {k.lower(): v for k, v in json.loads((TOOLS / "gear_stats.json").read_text()).items()}
art_vals = json.loads((TOOLS / "artifact_values.json").read_text())
game = json.loads((TOOLS / "game_en.json").read_text(encoding="utf-8"))

RESKIN = re.compile(r"Spooky|Winter|Year\d|Spring|Anniv", re.I)
NOISE_VAL = re.compile(r"CameraShake|Timeline|Oscillation|Template|Shake|PhysicsAsset|AggGeom|TypeHierarchy|Weight =|SharedPassiveCooldown|DropCone|Item Drop Effect|Thresholds", re.I)
item_ids = {v: k.split("/")[-1] for k, v in game.items()
            if k.startswith("ItemType/") and not re.match(r"(Flavour|Desc)_", k.split("/")[-1])}
descs = {k.split("/")[-1][5:].lower(): v for k, v in game.items() if k.startswith("ItemType/Desc_")}

# Artifacts whose folder isn't named after their item id
FOLDERS = {"Corrupted Seeds": "CorruptedSeeds", "Satchel of Elements": "SatchelOfTheElements", "Buzzy Nest": "Beenest",
           "Enchanter's Tome": "EnchantersTome", "Tome of Duplication": "TomeOfDuplication",
           "Satchel of Elixirs": "SatchelofNeed", "Satchel of Snacks": "SatchelofNourishment"}
F = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill("solid", fgColor="1F3864")
INPUT = PatternFill("solid", fgColor="FFF2CC")
UNIQUE = PatternFill("solid", fgColor="FCE4D6")
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="top", horizontal="center")


def item_id(x):
    return item_ids.get(x["name"]) or x["id"]


def items(cat):
    out = []
    for x in wiki:
        if x["category"] != cat or x["section"] == "Unused":
            continue
        iid = item_id(x)
        if RESKIN.search(iid or "") or x["name"] == "Corrupted Pumpkin":
            continue
        out.append({**x, "iid": iid})
    if cat != "Artifact":
        out.sort(key=lambda x: (x["group"].lower(), x["unique"], x["name"]))
    return out


def what_it_does(x):
    parts = list(x["properties"])
    if x["builtin"]:
        parts.append("Built-in: " + ", ".join(x["builtin"]))
    return "\n".join(f"- {p}" for p in parts) if parts else "-"


def fmt(v, nd=2):
    if v is None:
        return "-"
    if not isinstance(v, float):
        return str(v)
    t = f"{v:.{nd}f}"
    return t.rstrip("0").rstrip(".") if "." in t else t


def melee_cols(x):
    s = stats.get(x["iid"].lower()) or stats.get(x["id"].lower())
    hits = [h for h in (s or {}).get("hits", []) if h["damage"]]
    if not hits:
        return ["-", "-", "-"]
    shown = hits if len(hits) <= 6 else hits[:5]
    more = "" if len(hits) <= 6 else f" ... ({len(hits)} hits)"
    dmg = " / ".join(fmt(h["damage"], 0) for h in shown) + more
    tim = " / ".join(fmt(h["time"]) for h in shown) + more
    total_t = sum((h["time"] or 0) + (h["cooldown"] or 0) for h in hits)
    dps = round(sum(h["damage"] for h in hits) / total_t) if total_t else "-"
    if any((h.get("interval") or 0) > 0 or (h.get("loop") or 1) > 1 for h in hits):
        dps = "continuous / looping"
    return [dmg, tim, dps]


def ranged_cols(x):
    s = stats.get(x["iid"].lower()) or stats.get(x["id"].lower()) or {}
    if not s:
        return ["-"] * 5
    return [fmt(s.get("damage_factor")) if s.get("damage_factor") is not None else "1 (default)",
            fmt(s.get("charge_mult")) + (f" after {fmt(s.get('charge_time'))} s" if s.get("charge_time") else ""),
            fmt(s.get("attack_rate")) if s.get("attack_rate") is not None else "default",
            fmt(s.get("projectiles")), fmt(s.get("ammo")) if s.get("ammo") is not None else "default"]


def artifact_values(x):
    folder = FOLDERS.get(x["name"], x["iid"])
    vals = [v for v in art_vals.get(folder, []) if not NOISE_VAL.search(v)]
    return vals


def header(ws, row, headers, widths):
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row, c, h)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = HDR
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w


def put(ws, r, values, inputs=(), center=(), bold=(), unique=False):
    for c, v in enumerate(values, 1):
        cell = ws.cell(r, c, v)
        cell.font = Font(name=F, size=10, bold=c in bold)
        cell.alignment = CENTER if c in center else WRAP
        cell.border = BORDER
        if c in inputs:
            cell.fill = INPUT
        elif unique and c in bold:
            cell.fill = UNIQUE


# ---------------------------------------------------------------------------------------- detailed
def detailed():
    wb = Workbook()
    specs = {
        "Melee": (["#", "Type", "Name", "Unique", "What it does (in-game)", "Damage per hit (combo)",
                   "Swing time per hit (s)", "Approx. damage / sec", "Editable", "Verdict", "Your change"],
                  [5, 16, 24, 8, 46, 24, 22, 11, 9, 11, 36], melee_cols),
        "Ranged": (["#", "Type", "Name", "Unique", "What it does (in-game)", "Arrow damage factor",
                    "Charged shot (x damage after charge)", "Seconds between shots", "Arrows per shot",
                    "Arrow stack", "Editable", "Verdict", "Your change"],
                   [5, 16, 24, 8, 46, 11, 16, 11, 9, 9, 9, 11, 36], ranged_cols),
        "Armor": (["#", "Type", "Name", "Unique", "What it does (in-game)", "Editable", "Verdict", "Your change"],
                  [5, 18, 26, 8, 60, 9, 11, 40], None),
        "Artifacts": (["#", "Group", "Name", "What it does (in-game)", "Cooldown / souls", "Values in game files",
                       "Editable", "Verdict", "Your change"],
                      [5, 18, 22, 52, 22, 40, 9, 11, 36], None),
    }
    notes = {
        "Melee": "Editable: pak = per-hit damage, swing time, range, knockback and stun are in the game files. "
                 "Built-in enchantments are hard-coded (code patch). Approx. damage / sec = combo damage / combo time "
                 "(ignores area, knockback and animation cancelling). Damage is relative (power scales it).",
        "Ranged": "Editable: pak = damage factor, charge, fire rate, arrows per shot and arrow stack are in the game "
                  "files. Built-in enchantments are hard-coded. 'default' = engine default, not stored per item.",
        "Armor": "Editable: native = armor properties (damage reduction, cooldowns, auras...) are hard-coded in the "
                 "game's code. Changing them needs a code patch (research needed).",
        "Artifacts": "Editable: partial = some values are in the game files (listed); cooldowns and most effects "
                     "are hard-coded (code patch). native = nothing useful in the game files.",
    }
    dv_verdict = '"Keep,Buff,Nerf,Rework,Skip"'
    cats = {"Melee": "Melee", "Ranged": "Ranged", "Armor": "Armor", "Artifacts": "Artifact"}
    first = True
    for title, (headers, widths, colfn) in specs.items():
        ws = wb.active if first else wb.create_sheet(title)
        ws.title, first = title, False
        ws["A1"] = f"Sequel Mod - {title} Rebalance"
        ws["A1"].font = Font(name=F, size=14, bold=True)
        ws["A2"] = notes[title]
        ws["A2"].font = Font(name=F, size=9, italic=True)
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
        ws.row_dimensions[2].height = 30
        header(ws, 4, headers, widths)
        dv = DataValidation(type="list", formula1=dv_verdict, allow_blank=True)
        ws.add_data_validation(dv)
        r = 5
        for i, x in enumerate(items(cats[title]), 1):
            if title in ("Melee", "Ranged"):
                row = [i, x["group"], x["name"], "Yes" if x["unique"] else "", what_it_does(x)] + colfn(x) + ["pak", "", ""]
            elif title == "Armor":
                row = [i, x["group"], x["name"], "Yes" if x["unique"] else "", what_it_does(x), "native", "", ""]
            else:
                vals = artifact_values(x)
                desc = descs.get(x["iid"].lower(), "")
                does = "\n".join(f"- {p}" for p in x["properties"] if p)
                text = (desc + "\n" + does).strip() if desc and desc.lower() not in does.lower() else does
                row = [i, x["section"].replace(" artifacts", "").replace("Defensive, healing, and cooperative", "Defensive / healing"),
                       x["name"], text, x["cooldown"] or "-", "\n".join(vals) or "-", "partial" if vals else "native", "", ""]
            n = len(row)
            put(ws, r, row, inputs=(n - 1, n), center=(1, 4) if title != "Artifacts" else (1,), bold=(3,),
                unique=bool(x.get("unique")))
            dv.add(ws.cell(r, n - 1).coordinate)
            r += 1
        ws.freeze_panes = ws.cell(5, 4)
        ws.auto_filter.ref = f"A4:{ws.cell(r - 1, len(headers)).coordinate}"
    wb.save(ROOT / "Gear_Rebalance.xlsx")


# ---------------------------------------------------------------------------------------- simple
def simple():
    wb = Workbook()
    first = True
    for title, cat in (("Melee", "Melee"), ("Ranged", "Ranged"), ("Armor", "Armor"), ("Artifacts", "Artifact")):
        ws = wb.active if first else wb.create_sheet(title)
        ws.title, first = title, False
        ws["A1"] = f"Minecraft Dungeons {title}"
        ws["A1"].font = Font(name=F, size=14, bold=True)
        group = "Group" if cat == "Artifact" else "Type"
        header(ws, 3, ["Name", group, "What it does in-game", "Tier", "Proposed changes"], [26, 18, 64, 8, 45])
        dv = DataValidation(type="list", formula1='"S,A,B,C,D,F"', allow_blank=True)
        ws.add_data_validation(dv)
        r = 4
        for x in items(cat):
            if cat == "Artifact":
                grp = x["section"].replace(" artifacts", "").replace("Defensive, healing, and cooperative", "Defensive / healing")
                text = "\n".join(p for p in x["properties"] if p) + (f"\nCooldown: {x['cooldown']}" if x["cooldown"] else "")
            else:
                grp = x["group"] + (" (unique)" if x["unique"] else "")
                text = what_it_does(x).replace("- ", "", 1).replace("\n- ", "\n")
            put(ws, r, [x["name"], grp, text, "", ""], center=(4,), bold=(1,))
            dv.add(f"D{r}")
            r += 1
        ws.cell(r + 1, 1, "Item effects from the Minecraft Wiki").font = Font(name=F, size=8, italic=True, color="808080")
        ws.freeze_panes = "B4"
    wb.save(ROOT / "Gear_Tier_List.xlsx")


if __name__ == "__main__":
    detailed()
    simple()
    for c in ("Melee", "Ranged", "Armor", "Artifact"):
        print(c, len(items(c)))
