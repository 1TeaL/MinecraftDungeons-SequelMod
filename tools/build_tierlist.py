"""Builds Enchantment_Tier_List.xlsx (simple sheet for friends) from tools/wiki_enchantments.json."""
import json
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SRC, OUT = sys.argv[1:3]
rows = json.load(open(SRC, encoding="utf-8"))

SPECIAL = {
    "Snowball": "Triggers every 5 / 3 / 1 seconds",
    "Pain Cycle": "Max stacks deal 3x / 4x / 5x damage",
    "Shadow Blast": "Blast radius 2 / 3 / 4 blocks",
    "Beast Boss": "Pet damage +20% / +40% / +60%",
    "Exploding": "Deals 20% / 40% / 60% of enemy health as damage",
    "Piercing": "Every 3rd / 2nd / every arrow",
    "Electrified": "X / 2X / 3X damage",
    "Artifact Charge": "Next 1 / 2 / 3 attacks charged",
}
DESC = {"Rushdown": "100% increased speed temporarily after defeating a mob."}


def values_text(r):
    if r["name"] in SPECIAL:
        return SPECIAL[r["name"]]
    t = [v.strip() for v in r["t"]]
    if not any(t):
        return ""
    tt = re.sub(r"^[#%]\s*", "", r["tiertext"].strip()).rstrip(".").strip()
    joined = " / ".join(t)
    if re.search(r"\bX\b|\bx\b", tt) and not any("X" in v for v in t):
        return re.sub(r"\bX\b|\bx\b", joined, tt, count=1)
    return f"{joined} {tt}".strip()


F = "Arial"
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
SECTION = {"Melee": "FCE4D6", "Ranged": "DDEBF7", "Armor": "E2EFDA", "Unique-only": "EDE7F6"}

wb = Workbook()
ws = wb.active
ws.title = "Enchantments"
ws["A1"] = "Minecraft Dungeons Enchantments"
ws["A1"].font = Font(name=F, size=14, bold=True)

headers = ["Enchantment", "What it does (Tier I / II / III)", "Tier", "Proposed changes"]
for c, (h, w) in enumerate(zip(headers, [22, 70, 8, 45]), 1):
    cell = ws.cell(3, c, h)
    cell.font = Font(name=F, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F3864")
    cell.border = BORDER
    cell.alignment = Alignment(vertical="center")
    ws.column_dimensions[cell.column_letter].width = w

dv = DataValidation(type="list", formula1='"S,A,B,C,D,F"', allow_blank=True)
ws.add_data_validation(dv)

r = 4
for cat in SECTION:
    items = sorted((x for x in rows if x["cat"] == cat), key=lambda x: x["name"])
    ws.cell(r, 1, cat).font = Font(name=F, bold=True)
    for c in range(1, 5):
        ws.cell(r, c).fill = PatternFill("solid", fgColor=SECTION[cat])
        ws.cell(r, c).border = BORDER
    r += 1
    for x in items:
        desc = DESC.get(x["name"], x["desc"].strip())
        vals = values_text(x)
        text = desc + ("\n" + vals if vals else "")
        for c, v in enumerate([x["name"], text, "", ""], 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name=F, size=10, bold=(c == 1))
            cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c == 3 else None)
            cell.border = BORDER
        dv.add(f"C{r}")
        r += 1

ws.cell(r + 1, 1, "X = base damage. Tier values from the Minecraft Wiki").font = Font(name=F, size=8, italic=True, color="808080")
ws.freeze_panes = "B4"
wb.save(OUT)
print("rows:", len(rows))
