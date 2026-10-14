"""Builds Enchantment_Rebalance.xlsx from the game's localization + blueprint dump."""
import json
import re
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

LOC, DUMP, OUT = sys.argv[1:4]
loc = json.load(open(LOC, encoding="utf-8"))

# --- blueprint values from enchant_dump.txt ---------------------------------
SKIP = re.compile(r"^(UberGraphFrame|UIData|TypeId|TreatAsInherrent|bAlwaysTrigger|PrimaryComponentTick|"
                  r"bStartWithTickEnabled|.*Sound|.*Class|Effect|.*Effect|EffectsForLevel|EffectPerLevel|"
                  r"ChargingEffects|SummonHelper|CriticalDamageType|SmitingDamageType|WeakeningDamageType|"
                  r"ExplosionTag|DeniedProjectiles|ProjectileToDeflectedClasses|EffectMap|PoisonCloud|"
                  r"ActorToSpawn|Level|enablePushback|RemoteRole|Overlap|Movement|Root.*|DamageComponent)\b")
values = {}
cur = None
for line in open(DUMP, encoding="utf-8"):
    if line.startswith("=== "):
        cur = line[4:].strip()[3:]  # strip BP_
        values[cur] = []
    elif cur and re.match(r"^ {4,}\S", line) and " = " in line.strip():
        k, _, v = line.strip().partition(" = ")
        if not SKIP.match(k.split("[")[0]):
            if re.fullmatch(r"\[[0-9., e-]+\]", v):
                v = "[" + ", ".join(f"{float(x):g}" for x in v.strip("[]").split(",")) + "]"
            values[cur].append(f"{k} = {v}")
    elif cur and line.strip().endswith("):"):
        pass

# --- curated enchantment list ------------------------------------------------
# (category, locres key, blueprint name, note)
M, R, A = "Melee", "Ranged", "Armor"
ROWS = [
    (M, "Backstabber", "Backstabber", ""),
    (M, "Anima", "AnimaConduitMelee", ""),
    (M, "DamageSynergy", "DamageSynergy", ""),
    (M, "BusyBee", "BusyBee", "Weak: bees die fast and are small"),
    (M, "Chains", "Chains", ""),
    (M, "Committed", "Committed", ""),
    (M, "Critical", "CriticalHit", ""),
    (M, "DynamoMelee", "DynamoMelee", ""),
    (M, "Echo", "Echo", ""),
    (M, "EnigmaMelee", "EnigmaResonatorMelee", ""),
    (M, "Exploding", "Exploding", ""),
    (M, "Fire", "FireAspect", "Weak: low burn damage"),
    (M, "Freezing", "Freezing", ""),
    (M, "GravityMelee", "GravityMelee", ""),
    (M, "GuardingStrike", "GuardingStrike", "Weak: short shield"),
    (M, "BaneOfIllagers", "BaneOfIllagers", "Niche: illagers only"),
    (M, "JunglePoisonMelee", "JunglePoisonMelee", ""),
    (M, "Leeching", "Leeching", ""),
    (M, "Looting", "Looting", "Weak: rarely worth a slot"),
    (M, "PainCycle", "PainCycle", ""),
    (M, "Poisoned", "PoisonedMelee", "Weak: slow, small cloud"),
    (M, "Prospector", "Prospector", "Weak: emeralds rarely matter"),
    (M, "RadianceMelee", "RadianceMelee", ""),
    (M, "Rampaging", "Rampaging", "Weak: low trigger chance"),
    (M, "PotionThirstMelee", "PotionThirstMelee", ""),
    (M, "Rushdown", "Rushdown", "Weak: +10% speed for 0.8s"),
    (M, "Sharpness", "Sharpness", ""),
    (M, "Shock", "Shockwave", ""),
    (M, "Smiting", "Smiting", "Niche: undead only"),
    (M, "Soul", "SoulSiphon", "Weak: 20% chance, 1 soul"),
    (M, "Stunning", "Stunning", ""),
    (M, "Swirling", "Swirling", ""),
    (M, "Thundering", "Thundering", ""),
    (M, "VoidStrikeMelee", "VoidTouchedMelee", ""),
    (M, "Weakening", "Weakening", "Weak: small range and effect"),
    (R, "Accelerate", "Accelerating", ""),
    (R, "AnimaRanged", "AnimaConduitRanged", ""),
    (R, "BonusShot", "BonusShot", "Weak: extra arrow does little damage"),
    (R, "ChainReaction", "ChainReaction", ""),
    (R, "CooldownShot", "CooldownShot", ""),
    (R, "DynamoRanged", "DynamoRanged", ""),
    (R, "EnigmaRanged", "EnigmaResonatorRanged", ""),
    (R, "FuseShot", "FuseShot", "Weak: slow, unreliable"),
    (R, "GravityRanged", "Gravity", ""),
    (R, "Growing", "Growing", ""),
    (R, "Infinity", "Infinity", ""),
    (R, "LevitationShot", "levitationShot", ""),
    (R, "Multi", "MultiShot", ""),
    (R, "MultiCharge", "MultiCharge", ""),
    (R, "Piercing", "Piercing", ""),
    (R, "JunglePoison", "JunglePoisonRanged", ""),
    (R, "Poisoned", "PoisonedRanged", "Weak: slow, small cloud"),
    (R, "Power", "Power", ""),
    (R, "Punch", "Punch", "Weak: knockback often hurts you"),
    (R, "RadianceRanged", "RadianceRanged", ""),
    (R, "RapidFire", "RapidFire", ""),
    (R, "PotionThirstRanged", "PotionThirstRanged", ""),
    (R, "Ricochet", "Ricochet", ""),
    (R, "ShadowBarbRanged", "ShadowBarbRanged", ""),
    (R, "ShadowShot", "ShadowShot", ""),
    (R, "ShockWeb", "ShockWeb", ""),
    (R, "Supercharge", "Supercharge", "Weak: small bonus, knockback"),
    (R, "Tempo", "TempoTheft", ""),
    (R, "Unchanting", "Unchanting", ""),
    (R, "VoidStrikeRanged", "VoidTouchedRanged", ""),
    (R, "WildRage", "WildRage", "Weak: enraged mobs are a mixed bag"),
    (A, "Acrobat", "Acrobat", ""),
    (A, "BagOfSouls", "BagOfSouls", ""),
    (A, "BeastBoss", "BeastBoss", ""),
    (A, "BeastBurst", "BeastBurst", ""),
    (A, "BeastSurge", "BeastSurge", ""),
    (A, "Burning", "Burning", "Weak: low damage"),
    (A, "BurstBowstring", "BurstBowstring", ""),
    (A, "Chilling", "Chilling", "Weak: tiny effect"),
    (A, "Celerity", "Celerity", ""),
    (A, "Cowardice", "Cowardice", ""),
    (A, "DeathBarter", "DeathBarter", ""),
    (A, "Deflect", "Deflecting", ""),
    (A, "DippingPoison", "DippingPoison", ""),
    (A, "Electrified", "Electrified", ""),
    (A, "Explorer", "Explorer", "Weak: tiny heal"),
    (A, "FinalShout", "FinalShout", ""),
    (A, "FireFocus", "FireFocus", ""),
    (A, "FireTrail", "FireTrail", "Weak: low damage"),
    (A, "FoodReserves", "FoodReserves", "Weak: food is barely useful"),
    (A, "Frenzied", "Frenzied", ""),
    (A, "GravityPulse", "GravityPulse", ""),
    (A, "HealthSynergy", "HealthSynergy", ""),
    (A, "ResurrectionSurge", "ResurrectionSurge", ""),
    (A, "LightningFocus", "LightningFocus", ""),
    (A, "EmeraldDivination", "EmeraldDivination", "Weak: emeralds rarely matter"),
    (A, "LuckOfTheSea", "LuckOfTheSea", ""),
    (A, "MultiDodge", "MultiDodge", ""),
    (A, "PoisonFocus", "PoisonFocus", ""),
    (A, "PotionFortification", "PotionFortification", ""),
    (A, "Reckless", "Reckless", ""),
    (A, "Recycler", "Recycler", "Weak: needs lots of hits"),
    (A, "RollCharge", "RollCharge", ""),
    (A, "Flee", "Flee", ""),
    (A, "Snowing", "Snowing", "Weak: low damage"),
    (A, "SoulFocus", "SoulFocus", "Weaker than other Focus (+10% vs +25%)"),
    (A, "SpiritSpeed", "SpiritSpeed", ""),
    (A, "SpeedSynergy", "SpeedSynergy", ""),
    (A, "SurpriseGift", "SurpriseGift", ""),
    (A, "Swiftfooted", "Swiftfooted", ""),
    (A, "Thorns", "Thorns", "Weak: low damage"),
    (A, "ThriveUnderPressure", "ThriveUnderPressure", ""),
    (A, "TumbleBee", "TumbleBee", ""),
    (A, "ShadowFeast", "ShadowFeast", "Category unsure"),
    (A, "ShadowFlash", "ShadowFlash", "Category unsure"),
    (R, "ArtifactCharge", "ArtifactCharge", ""),
]

# Unique-item / mob / internal enchantments (not on the normal enchant pool)
OTHER = ["AncientGuardianThorns", "Barrier", "BowsBoon", "CaveSpiderPoisonEnchantment",
         "ChargingAcceleration", "CogCrossbowEnchantment", "CurrentImmunity", "DamageCounter",
         "DeflectingWithShield", "FreezingAoe", "FreezingRanged", "HeavyweightEnchantment", "Huge",
         "HuntingBowEnchantment", "HuntingBowTaggedEnchantment", "Invisible", "Knockback",
         "MultiCharge", "PassiveRegen", "ProjectileCounter", "PushVoumeImmunity", "ReliableRicochet",
         "ResurrectSurroundingMobs", "SlowBowEnchantment", "SlowImmunity", "SlowResistanceEnchantment",
         "SpongeStrike", "Altruistic", "UnderwaterImmunity", "VoidStrikeImmunity", "WindBowEnchantment",
         "WindImmunity", "WindResistance", "WitherEnchantmentMelee", "WitherEnchantmentRanged",
         "BlindCloud", "Regeneration", "Shielding", "Protection", "Quick", "FastAttack",
         "DoubleDamage", "Leeching", "SharedPain", "VoidBlast", "ShulkerSentry"]
OTHER = [o for o in OTHER if o not in ("MultiCharge", "Leeching")]


def name(key):
    n = loc.get(f"Enchantment/{key}", "")
    return n or {"Celerity": "Cool Down", "AnimaRanged": "Anima Conduit", "Soul": "Soul Siphon",
                 "Accelerate": "Accelerate"}.get(key, key)


MY_DESC = {
    "EnigmaMelee": loc.get("Enchantment/Enigma_desc", ""),
    "Accelerate": loc.get("Enchantment/Accelerating_desc", ""),
    "Deflect": loc.get("Enchantment/Deflecting_desc", ""),
    "EmeraldDivination": loc.get("Enchantment/EmeraldDivination_effect_desc", ""),
    "Flee": loc.get("Enchantment/Flee_effect_desc", ""),
    "DeathBarter": loc.get("Enchantment/DeathBarter_effect_desc", ""),
    "ChainReaction": "Chance to fire arrows in all directions when an arrow hits a mob.",
    "JunglePoisonMelee": "Hits poison the target, dealing damage over time.",
    "JunglePoison": "Arrows poison the target, dealing damage over time.",
    "Reckless": "Increases your melee damage but lowers your max health.",
}

# Only a minor value is in the files; the core number (damage, chance) is native.
PARTIAL = {"Thorns", "Electrified", "FireTrail", "PoisonedMelee", "PoisonedRanged",
           "RadianceMelee", "Thundering", "BeastBurst"}


def desc(key):
    d = loc.get(f"Enchantment/{key}_desc", "")
    return d if d and not d.isupper() else MY_DESC.get(key, d)


def effect(key):
    e = loc.get(f"Enchantment/{key}_effect", "")
    return "" if e.isupper() else e


# --- workbook -----------------------------------------------------------------
F = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
CAT_FILL = {M: "FCE4D6", R: "DDEBF7", A: "E2EFDA"}
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()
ws = wb.active
ws.title = "Enchantments"

ws["A1"] = "Sequel Mod - Enchantment Rebalance"
ws["A1"].font = Font(name=F, size=14, bold=True)
ws["A2"] = "Editable: pak = in game files | partial = only a minor value in game files | native = hard-coded in the exe"
ws["A2"].font = Font(name=F, size=9, italic=True)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:I2")
ws.row_dimensions[2].height = 18

headers = ["#", "Type", "Enchantment", "What it does (in-game)", "Scales per level (tooltip)",
           "Values in game files", "Editable", "Verdict", "Your change"]
widths = [5, 9, 22, 48, 26, 36, 10, 12, 40]
HR = 4
for c, (h, w) in enumerate(zip(headers, widths), 1):
    cell = ws.cell(HR, c, h)
    cell.font = Font(name=F, bold=True, color="FFFFFF")
    cell.fill = HDR_FILL
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border = BORDER
    ws.column_dimensions[cell.column_letter].width = w

dv = DataValidation(type="list", formula1='"Keep,Buff,Nerf,Rework,Skip"', allow_blank=True)
ws.add_data_validation(dv)

r = HR + 1
for i, (cat, key, bp, note) in enumerate(ROWS, 1):
    vals = values.get(bp, [])
    row = [i, cat, name(key), desc(key), effect(key), "\n".join(vals) or "-",
           ("partial" if bp in PARTIAL else "pak") if vals else "native", "", ""]
    for c, v in enumerate(row, 1):
        cell = ws.cell(r, c, v)
        cell.font = Font(name=F, size=10, bold=(c == 3))
        cell.alignment = WRAP
        cell.border = BORDER
        if c == 2:
            cell.fill = PatternFill("solid", fgColor=CAT_FILL[cat])
        if c in (8, 9):
            cell.fill = INPUT_FILL
    dv.add(f"H{r}")
    r += 1

ws.freeze_panes = ws.cell(HR + 1, 4)
ws.auto_filter.ref = f"A{HR}:I{r - 1}"

# Other tab
ws2 = wb.create_sheet("Unique & internal")
ws2["A1"] = "Unique items, mobs and internal effects"
ws2["A1"].font = Font(name=F, size=9, italic=True)
ws2.merge_cells("A1:E1")
for c, (h, w) in enumerate(zip(["Blueprint", "Values in game files", "Editable", "Verdict", "Your change"],
                               [30, 45, 10, 12, 40]), 1):
    cell = ws2.cell(3, c, h)
    cell.font = Font(name=F, bold=True, color="FFFFFF"); cell.fill = HDR_FILL; cell.border = BORDER
    ws2.column_dimensions[cell.column_letter].width = w
dv2 = DataValidation(type="list", formula1='"Keep,Buff,Nerf,Rework,Skip"', allow_blank=True)
ws2.add_data_validation(dv2)
for j, bp in enumerate(sorted(OTHER), 4):
    vals = values.get(bp, [])
    for c, v in enumerate([bp, "\n".join(vals) or "-", "pak" if vals else "native", "", ""], 1):
        cell = ws2.cell(j, c, v)
        cell.font = Font(name=F, size=10); cell.alignment = WRAP; cell.border = BORDER
        if c >= 4:
            cell.fill = INPUT_FILL
    dv2.add(f"D{j}")
ws2.freeze_panes = "A4"

wb.save(OUT)

notes = ["# Enchantment notes", "", "Notes that go with Enchantment_Rebalance.xlsx.", "",
         "## Notes per enchantment", ""]
for cat, key, bp, note in ROWS:
    if note:
        notes.append(f"- **{name(key)}** ({cat}): {note}")
notes += ["", "## Descriptions I wrote", "",
          "These had no description in the game text, so the sheet uses my own wording:", ""]
notes += [f"- {name(k)} ({'Melee' if k == 'JunglePoisonMelee' else 'Ranged' if k in ('JunglePoison', 'ChainReaction') else 'Armor'})"
          for k in ("ChainReaction", "JunglePoisonMelee", "JunglePoison", "Reckless")]
notes += ["", "## Unique & internal tab", "",
          "- Shared Pain, Void Blast and Shulker Sentry only have placeholder text in the game files; they look like cut or unfinished enchantments.",
          "", "## Per-level values", "",
          "- The game files don't store per-level numbers for most enchantments. The 'Scales per level' column is only the tooltip template; real per-level values for native enchantments need UE4SS to read."]
NL = chr(10)
open(OUT.replace("Rebalance.xlsx", "Notes.md"), "w", encoding="utf-8").write(NL.join(notes) + NL)
print("rows:", len(ROWS), "other:", len(OTHER), "pak:", sum(1 for x in ROWS if values.get(x[2])))
missing = [x[1] for x in ROWS if not desc(x[1])]
print("no desc:", missing)
print("no bp:", [x[2] for x in ROWS if x[2] not in values])
